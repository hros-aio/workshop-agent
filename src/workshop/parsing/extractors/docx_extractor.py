"""Word (.docx) extraction (python-docx)."""

from __future__ import annotations

import io
import uuid

from docx import Document as DocxDocument

from workshop.parsing.extractors.base import ExtractorError
from workshop.parsing.models import ExtractedFact, SourceRef

_DOCX_MIME = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


class DocxExtractor:
    def supports(self, mime_type: str) -> bool:
        return mime_type in _DOCX_MIME

    def extract(self, data: bytes, *, file_name: str, document_version_id: uuid.UUID) -> list[ExtractedFact]:
        facts: list[ExtractedFact] = []
        try:
            document = DocxDocument(io.BytesIO(data))
        except Exception as exc:
            raise ExtractorError(f"Failed to parse DOCX: {exc}") from exc

        for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                facts.append(
                    ExtractedFact(
                        text=text,
                        source_ref=SourceRef(
                            document_version_id=document_version_id,
                            file_name=file_name,
                            row=paragraph_index,
                            excerpt=text[:500],
                        ),
                    )
                )

        for table_index, table in enumerate(document.tables, start=1):
            sheet = f"Table{table_index}"
            for row_index, row in enumerate(table.rows, start=1):
                cells = [cell.text.strip() for cell in row.cells]
                text = " | ".join(c for c in cells if c)
                if text:
                    facts.append(
                        ExtractedFact(
                            text=text,
                            source_ref=SourceRef(
                                document_version_id=document_version_id,
                                file_name=file_name,
                                sheet=sheet,
                                row=row_index,
                                excerpt=text[:500],
                            ),
                        )
                    )
        return facts
